"""Async MOEX history loader with incremental updates.

Основные улучшения:
 - Асинхронные HTTP-запросы (aiohttp) с ограничением параллелизма.
 - Инкрементальная догрузка: старт от последней даты в БД +1 день либо от START_DATE.
 - Минимизируется повторное форматирование/парсинг дат.
 - Логирование вместо print.
 - Опциональный экспорт в Excel / JSON.
 - Параметры через CLI (шаг диапазона, максимальная конкуренция, отключение экспорта и т.п.).
"""

from __future__ import annotations

import argparse
import asyncio
import json
import logging
import os
import sqlite3
import datetime as dt
from typing import Sequence, Iterable, Any

import aiohttp
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# --- Константы / настройки (можно переопределять переменными окружения) ---
DB_PATH = os.getenv("MOEX_DB_PATH", "moex_data.db")
START_DATE = dt.date(2022, 1, 1)
DEFAULT_BOARD = "TQBR"
DATE_STEP_DAYS = int(os.getenv("MOEX_DATE_STEP", "100"))
MAX_CONCURRENCY = int(os.getenv("MOEX_MAX_CONCURRENCY", "8"))
HTTP_TIMEOUT = int(os.getenv("MOEX_HTTP_TIMEOUT", "20"))
HTTP_RETRIES = int(os.getenv("MOEX_HTTP_RETRIES", "3"))
HTTP_BACKOFF = float(os.getenv("MOEX_HTTP_BACKOFF", "0.5"))
RATE_LIMIT_RPS = float(os.getenv("MOEX_RATE_LIMIT", "0"))  # 0 = disabled
class RateLimiter:
    """Simple async token bucket rate limiter.

    rate: tokens (requests) per second.
    burst: maximum accumulated tokens.
    """

    def __init__(self, rate: float, burst: int | None = None):
        if rate <= 0:
            raise ValueError("rate must be > 0")
        self.rate = rate
        self.capacity = burst if burst and burst > 0 else max(1, int(rate))
        self.tokens = self.capacity
        self.updated = asyncio.get_event_loop().time()
        self._lock = asyncio.Lock()

    async def acquire(self) -> None:
        async with self._lock:
            now = asyncio.get_event_loop().time()
            # Refill tokens
            elapsed = now - self.updated
            if elapsed > 0:
                refill = elapsed * self.rate
                if refill > 0:
                    self.tokens = min(self.capacity, self.tokens + refill)
                    self.updated = now
            if self.tokens >= 1:
                self.tokens -= 1
                return
            # Need to wait: compute time to next token
            needed = 1 - self.tokens
            wait_time = needed / self.rate
        # Release lock before sleeping
        await asyncio.sleep(wait_time)
        # Recursive attempt (could also loop)
        await self.acquire()



# --- Логирование ---
def setup_logging(level: str = "INFO") -> None:
    logging.basicConfig(
        level=level,
        format="%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )


# --- Утилиты дат ---
def get_date_ranges(start: dt.date, end: dt.date, step: int = DATE_STEP_DAYS) -> list[tuple[dt.date, dt.date]]:
    ranges: list[tuple[dt.date, dt.date]] = []
    current = start
    while current <= end:
        till = min(current + dt.timedelta(days=step - 1), end)
        ranges.append((current, till))
        current = till + dt.timedelta(days=1)
    return ranges


# --- Работа с БД ---
def get_last_dates(conn: sqlite3.Connection, secids: Iterable[str]) -> dict[tuple[str, str], str | None]:
    """Вернуть {(boardid, secid): max_trade_date or None}."""
    result: dict[tuple[str, str], str | None] = {}
    for secid in secids:
        rows = conn.execute("SELECT DISTINCT BOARDID FROM moex WHERE SECID=?", (secid,)).fetchall()
        boardids = [r[0] for r in rows] or [DEFAULT_BOARD]
        for board in boardids:
            max_date = conn.execute(
                "SELECT MAX(TRADEDATE) FROM moex WHERE SECID=? AND BOARDID=?",
                (secid, board),
            ).fetchone()[0]
            result[(board, secid)] = max_date
    return result


def read_instruments(conn: sqlite3.Connection, only: list[str] | None = None) -> list[str]:
    df = pd.read_sql("SELECT * FROM share", conn)
    if "share" not in df.columns:
        raise RuntimeError("В таблице share отсутствует колонка 'share'")
    series = df["share"].dropna().astype(str).unique().tolist()
    if only:
        filt = {s.upper() for s in only}
        series = [s for s in series if s.upper() in filt]
    return series


# --- Асинхронный HTTP слой ---
async def fetch_range(
    session: aiohttp.ClientSession,
    secid: str,
    board: str,
    dr_start: dt.date,
    dr_end: dt.date,
    semaphore: asyncio.Semaphore,
    rate_limiter: RateLimiter | None,
) -> tuple[list[list[Any]], list[str] | None]:
    url = (
        "https://iss.moex.com/iss/history/engines/stock/markets/shares/boards/"
        f"{board}/securities/{secid}.json?from={dr_start}&till={dr_end}"
    )
    attempt = 0
    await semaphore.acquire()
    try:
        while True:
            attempt += 1
            try:
                if rate_limiter:
                    await rate_limiter.acquire()
                async with session.get(url, timeout=HTTP_TIMEOUT) as resp:
                    if resp.status >= 500:
                        raise aiohttp.ClientResponseError(
                            resp.request_info, resp.history, status=resp.status, message="server error"
                        )
                    if resp.status != 200:
                        text = await resp.text()
                        logging.warning(
                            "HTTP %s (%s %s) %s", resp.status, secid, f"{dr_start}:{dr_end}", text[:200]
                        )
                        return [], None
                    data = await resp.json()
                    history = data.get("history") or {}
                    columns = history.get("columns")
                    rows = history.get("data") or []
                    # --- Pagination handling ---
                    # MOEX ISS provides a cursor dataset 'history.cursor' with columns including TOTAL & PAGESIZE
                    cursor = data.get("history.cursor") or {}
                    cursor_cols = cursor.get("columns") or []
                    cursor_data = cursor.get("data") or []
                    total_rows_reported = None
                    page_size = None
                    if cursor_cols and cursor_data:
                        # Build mapping from cursor row (assume first row contains meta)
                        first_row = cursor_data[0]
                        mapping = {cursor_cols[i]: first_row[i] for i in range(min(len(cursor_cols), len(first_row)))}
                        # Common field names: TOTAL, PAGESIZE (sometimes uppercase)
                        total_rows_reported = mapping.get("TOTAL") or mapping.get("total")
                        page_size = mapping.get("PAGESIZE") or mapping.get("pagesize")
                    # If total > len(rows), fetch remaining pages sequentially using start=offset
                    if total_rows_reported and page_size and isinstance(total_rows_reported, int) and isinstance(page_size, int):
                        fetched = len(rows)
                        start = fetched
                        # Safety cap to prevent infinite loop in case of inconsistent cursor
                        max_expected = total_rows_reported
                        page_index = 1
                        while fetched < total_rows_reported and start < max_expected:
                            page_url = f"{url}&start={start}"
                            try:
                                if rate_limiter:
                                    await rate_limiter.acquire()
                                async with session.get(page_url, timeout=HTTP_TIMEOUT) as page_resp:
                                    if page_resp.status != 200:
                                        txt = await page_resp.text()
                                        logging.warning("Page fetch status %s for %s start=%s: %s", page_resp.status, secid, start, txt[:120])
                                        break
                                    page_json = await page_resp.json()
                                    page_history = page_json.get("history") or {}
                                    page_rows = page_history.get("data") or []
                                    if not page_rows:
                                        break
                                    rows.extend(page_rows)
                                    fetched += len(page_rows)
                                    start += len(page_rows)
                                    page_index += 1
                                    logging.debug("Pagination page fetched %s (rows cumulative %s/%s) %s", page_index, fetched, total_rows_reported, secid)
                                    # Guard: if page_rows less than page_size, likely last page
                                    if len(page_rows) < page_size:
                                        break
                            except (aiohttp.ClientError, asyncio.TimeoutError) as page_exc:
                                logging.warning("Ошибка пагинации %s start=%s: %s", secid, start, page_exc)
                                break
                    logging.debug(
                        "Fetched %s rows for %s %s %s-%s (reported total=%s)", len(rows), secid, board, dr_start, dr_end, total_rows_reported
                    )
                    return rows, columns
            except (aiohttp.ClientError, asyncio.TimeoutError) as exc:
                if attempt >= HTTP_RETRIES:
                    logging.error("Fail %s after %s attempts: %s", secid, attempt, exc)
                    return [], None
                backoff = HTTP_BACKOFF * (2 ** (attempt - 1))
                logging.warning(
                    "Retry %s (%s %s-%s) after error: %s (sleep %.2fs)", attempt, secid, dr_start, dr_end, exc, backoff
                )
                await asyncio.sleep(backoff)
    finally:
        semaphore.release()


async def fetch_security(
    session: aiohttp.ClientSession,
    secid: str,
    board: str,
    ranges: Sequence[tuple[dt.date, dt.date]],
    semaphore: asyncio.Semaphore,
    rate_limiter: RateLimiter | None,
) -> tuple[list[list[Any]], list[str] | None]:
    tasks = [fetch_range(session, secid, board, r[0], r[1], semaphore, rate_limiter) for r in ranges]
    results = await asyncio.gather(*tasks)
    all_rows: list[list[Any]] = []
    columns: list[str] | None = None
    for rows, cols in results:
        if rows:
            all_rows.extend(rows)
        if cols and not columns:
            columns = cols
    return all_rows, columns


def normalize_dates(df: pd.DataFrame, date_cols: Iterable[str]) -> pd.DataFrame:
    for c in date_cols:
        if c in df.columns:
            df[c] = pd.to_datetime(df[c], errors="coerce")
    return df


def format_ru_dates(df: pd.DataFrame, date_cols: Iterable[str]) -> pd.DataFrame:
    for c in date_cols:
        if c in df.columns:
            df[c] = pd.to_datetime(df[c], errors="coerce").dt.strftime("%d.%m.%Y")
    return df


def save_json(data: list[dict[str, Any]], path: str) -> None:
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    logging.info("JSON сохранен: %s", path)


def save_excel(df: pd.DataFrame, path: str) -> None:
    df.to_excel(path, index=False)
    wb = load_workbook(path)
    ws = wb.active
    tab = Table(displayName="MOEXTable", ref=ws.dimensions)
    style = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=True
    )
    tab.tableStyleInfo = style
    ws.add_table(tab)
    wb.save(path)
    logging.info("Excel сохранен: %s", path)


def append_to_sqlite(conn: sqlite3.Connection, df: pd.DataFrame) -> int:
    before = conn.execute("SELECT COUNT(*) FROM moex").fetchone()[0] if table_exists(conn, "moex") else 0
    df.to_sql("moex", conn, if_exists="append", index=False)
    after = conn.execute("SELECT COUNT(*) FROM moex").fetchone()[0]
    return after - before


def table_exists(conn: sqlite3.Connection, name: str) -> bool:
    cur = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=? LIMIT 1",
        (name,),
    )
    return cur.fetchone() is not None


def resolve_start_date(base_start: dt.date, args: argparse.Namespace, last_date_raw: str | None) -> dt.date:
    """Determine the effective start date for a security considering overrides.

    Precedence (highest wins):
    1. --since YYYY-MM-DD (explicit absolute date)
    2. --days N (start = end_date - N + 1)
    3. Incremental (last_date + 1)
    4. Global START_DATE constant
    """
    # If user forced absolute since
    if getattr(args, "since", None):
        try:
            return dt.datetime.strptime(args.since, "%Y-%m-%d").date()
        except ValueError:
            logging.error("Некорректный формат --since: %s (ожидается YYYY-MM-DD)", args.since)
    # days relative window
    if getattr(args, "days", None) is not None:
        try:
            window_days = int(args.days)
            if window_days > 0:
                return args._effective_end_date - dt.timedelta(days=window_days - 1)
        except (TypeError, ValueError):
            logging.error("Некорректное значение --days: %s", args.days)
    # incremental from last_date_raw
    if last_date_raw:
        parsed = None
        for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                parsed = dt.datetime.strptime(last_date_raw, fmt).date()
                break
            except ValueError:
                pass
        if parsed:
            return parsed + dt.timedelta(days=1)
    # fallback
    return base_start


async def async_run(args: argparse.Namespace) -> None:
    setup_logging(args.log_level)
    end_date = dt.date.today() if args.to_date is None else dt.datetime.strptime(args.to_date, "%Y-%m-%d").date()
    # сохранить для расчёта --days
    args._effective_end_date = end_date  # type: ignore[attr-defined]
    async with aiohttp.ClientSession(headers={"User-Agent": "moex-loader/1.0"}) as session:
        semaphore = asyncio.Semaphore(args.max_concurrency)
        rate_limiter = None
        if getattr(args, "rate_limit", None) and args.rate_limit > 0:
            try:
                rate_limiter = RateLimiter(args.rate_limit)
                logging.info("Rate limiting включен: %.2f rps", args.rate_limit)
            except ValueError as e:
                logging.error("Не удалось включить rate limit: %s", e)
        # Поддержка dry-run: если включен, не открываем БД для записи, только читаем список инструментов из нее (если нужно)
        conn: sqlite3.Connection | None = None
        try:
            if args.dry_run:
                # В dry-run читаем список инструментов и last_dates (если есть БД), но не вставляем данные
                if os.path.exists(DB_PATH):
                    conn = sqlite3.connect(DB_PATH)
                    instruments = read_instruments(conn, args.instruments)
                    last_dates = get_last_dates(conn, instruments) if instruments else {}
                else:
                    logging.warning("БД %s не найдена — dry-run будет выполнен только для списка инструментов из аргументов", DB_PATH)
                    instruments = args.instruments or []
                    last_dates = {}
            else:
                conn = sqlite3.connect(DB_PATH)
                instruments = read_instruments(conn, args.instruments)
                if not instruments:
                    logging.warning("Нет инструментов для загрузки.")
                    return
                if not table_exists(conn, "moex"):
                    logging.info("Таблица moex отсутствует — будет создана при первой вставке.")
                last_dates = get_last_dates(conn, instruments)

            if not instruments:
                logging.warning("Список инструментов пуст.")
                return

            all_export_rows: list[dict[str, Any]] = []
            export_columns: list[str] | None = None
            summary: list[dict[str, Any]] = []

            for secid in instruments:
                boards = {b for (b, s) in last_dates.keys() if s == secid} or {DEFAULT_BOARD}
                total_rows_sec = 0
                for board in boards:
                    last_date_raw = last_dates.get((board, secid))
                    start_date = resolve_start_date(START_DATE, args, last_date_raw)
                    if start_date > end_date:
                        logging.info("%s %s: нет новых дат (start>%s)", secid, board, end_date)
                        continue
                    ranges = get_date_ranges(start_date, end_date, args.step_days)
                    logging.info("%s %s: диапазонов %s (с %s по %s)%s", secid, board, len(ranges), start_date, end_date, " [dry-run]" if args.dry_run else "")
                    rows, cols = await fetch_security(session, secid, board, ranges, semaphore, rate_limiter)
                    if not rows:
                        logging.info("%s %s: новых строк нет", secid, board)
                        continue
                    if cols and not export_columns:
                        export_columns = cols
                    rows = [r for r in rows if r and r[0] == board]
                    if not rows:
                        continue
                    df = pd.DataFrame(rows, columns=export_columns)
                    if {"BOARDID", "TRADEDATE", "SECID"}.issubset(df.columns):
                        df.drop_duplicates(subset=["BOARDID", "TRADEDATE", "SECID"], inplace=True)
                    if args.dry_run:
                        inserted = len(df)
                        logging.info("%s %s: (dry-run) потенциально добавлено %s строк", secid, board, inserted)
                    else:
                        inserted = append_to_sqlite(conn, df) if conn else 0
                        logging.info("%s %s: добавлено %s строк", secid, board, inserted)
                    total_rows_sec += inserted
                    if args.export or args.export_json:
                        df_export = df.copy()
                        df_export = format_ru_dates(df_export, ["TRADEDATE", "TRADE_SESSION_DATE"])
                        all_export_rows.extend(df_export.to_dict(orient="records"))
                summary.append({"SECID": secid, "rows": total_rows_sec})
        finally:
            if conn:
                conn.close()

        # Экспорт вне контекста соединения
        if (args.export or args.export_json) and all_export_rows and export_columns:
            export_df = pd.DataFrame(all_export_rows, columns=export_columns)
            if args.export:
                save_excel(export_df, args.export)
            if args.export_json:
                save_json(all_export_rows, args.export_json)

        # Итоговая сводка
        total_rows = sum(item["rows"] for item in summary)
        logging.info("==== СВОДКА ====")
        for item in summary:
            logging.info("%s: добавлено %s", item["SECID"], item["rows"])
        logging.info("Всего инструментов: %s, всего новых строк: %s", len(summary), total_rows)


def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Async MOEX history loader")
    p.add_argument("--instruments", nargs="*", help="Список SECID для ограничения выборки (опционально)")
    p.add_argument("--to-date", help="Дата окончания (YYYY-MM-DD), по умолчанию сегодня", dest="to_date")
    p.add_argument("--step-days", type=int, default=DATE_STEP_DAYS, help="Размер диапазона дат в днях (батч)")
    p.add_argument("--max-concurrency", type=int, default=MAX_CONCURRENCY, dest="max_concurrency", help="Максимум одновременных запросов")
    p.add_argument("--since", help="Принудительная дата начала (YYYY-MM-DD), перекрывает инкремент и --days")
    p.add_argument("--days", type=int, help="Ограничить диапазон последними N днями (игнорируется если задан --since)")
    p.add_argument("--dry-run", action="store_true", help="Загрузить и сформировать экспорт без записи в БД")
    p.add_argument("--rate-limit", type=float, help="Ограничить частоту запросов (RPS), 0 или пропуск = без лимита")
    p.add_argument("--export", nargs="?", const="moex_data.xlsx", help="Экспорт в Excel (опционально имя файла)")
    p.add_argument("--export-json", nargs="?", const="moex_data.json", help="Экспорт в JSON (опционально имя файла)")
    p.add_argument("--log-level", default="INFO", help="Уровень логирования (DEBUG/INFO/WARNING/...)")
    return p


def main() -> None:
    args = build_arg_parser().parse_args()
    try:
        asyncio.run(async_run(args))
    except KeyboardInterrupt:
        logging.warning("Прервано пользователем")
    except Exception as exc:  # noqa: BLE001
        logging.exception("Необработанная ошибка: %s", exc)


if __name__ == "__main__":  # pragma: no cover
    main()

